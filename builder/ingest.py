"""Ingest a case folder into one brief text + a tool API file.     python3 -m builder.ingest <case_folder> <mock_url> <out_dir>

Order: BRIEF.md, README.md, then every other file (sorted). Each file is converted to text:
md/txt as is; docx (python-docx, pandoc as fallback; tables -> markdown); pdf (pdftotext, pypdf as fallback); xlsx/xlsm (openpyxl, every sheet -> markdown table);
csv -> markdown table; html -> text; json/yaml that is not an OpenAPI file -> as text.
A json/yaml file above TEXT_CAP characters is not pasted: the brief gets a structural summary
(keys, counts, the list of services, and two or three example entries) and a note that the full file is test data for the mock.
Images are recorded as drawings, not read.
Each file is dated (from its name, a header, or the text) and typed (mail, meeting note, spec, FAQ/web page, drawing, prompt sheet).
Headers "=== FILE: name ===" and "=== META: date=...; type=...; author=... ===" separate the files.
Tool API: GET <mock_url>/openapi.json, else an OpenAPI json/yaml file in the folder.
"""
from __future__ import annotations

import csv
import io
import json
import re
import shutil
import subprocess
import sys
import urllib.request
from pathlib import Path

FIRST = ("brief.md", "readme.md")
TEXT = {".md", ".txt", ".markdown", ".rst"}
SKIP = {".zip", ".gz", ".mp3", ".wav", ".mp4", ".ds_store"}
IMAGE = {".png", ".jpg", ".jpeg", ".gif", ".webp", ".svg"}
TEXT_CAP = 60_000
SUMMARY_MARK = "NOTE: this file is above the brief size cap."


def _md_table(rows) -> str:
    rows = [[("" if c is None else str(c)).replace("\n", " ").replace("|", "/").strip() for c in r] for r in rows]
    rows = [r for r in rows if any(r)]
    if not rows:
        return ""
    w = max(len(r) for r in rows)
    rows = [r + [""] * (w - len(r)) for r in rows]
    out = ["| " + " | ".join(rows[0]) + " |", "|" + "---|" * w]
    out += ["| " + " | ".join(r) + " |" for r in rows[1:]]
    return "\n".join(out)


def read_docx(p: Path) -> str:
    try:
        import docx  # noqa: PLC0415
        d = docx.Document(str(p))
        from docx.table import Table  # noqa: PLC0415
        from docx.text.paragraph import Paragraph  # noqa: PLC0415
        out = []
        for el in d.element.body.iterchildren():   # document order: paragraphs and tables interleaved
            if el.tag.endswith("}p"):
                para = Paragraph(el, d)
                t = para.text.strip()
                if t:
                    st = (para.style.name or "") if para.style is not None else ""
                    out.append(("#" * int(st[-1]) + " " + t) if st.startswith("Heading ") and st[-1].isdigit() else t)
            elif el.tag.endswith("}tbl"):
                out.append(_md_table([[c.text for c in r.cells] for r in Table(el, d).rows]))
        return "\n\n".join(x for x in out if x)
    except Exception:  # noqa: BLE001
        if shutil.which("pandoc"):
            return subprocess.run(["pandoc", "-t", "gfm", str(p)], capture_output=True, text=True, timeout=60).stdout
        raise


def read_pdf(p: Path) -> str:
    if shutil.which("pdftotext"):
        r = subprocess.run(["pdftotext", "-layout", str(p), "-"], capture_output=True, text=True, timeout=120)
        if r.stdout.strip():
            return r.stdout
    from pypdf import PdfReader  # noqa: PLC0415
    return "\n\n".join((pg.extract_text() or "") for pg in PdfReader(str(p)).pages)


def read_xlsx(p: Path) -> str:
    import openpyxl  # noqa: PLC0415
    wb = openpyxl.load_workbook(str(p), data_only=True, read_only=True)
    out = []
    for ws in wb.worksheets:
        t = _md_table(list(ws.iter_rows(values_only=True)))
        if t:
            out.append(f"### Sheet: {ws.title}\n\n{t}")
    return "\n\n".join(out)


def read_csv(p: Path) -> str:
    raw = p.read_text(encoding="utf-8-sig", errors="replace")
    try:
        dialect = csv.Sniffer().sniff(raw[:4000])
    except csv.Error:
        dialect = csv.excel
    return _md_table(list(csv.reader(io.StringIO(raw), dialect)))


def read_html(p: Path) -> str:
    s = p.read_text(encoding="utf-8", errors="replace")
    s = re.sub(r"(?is)<(script|style).*?</\1>", " ", s)
    s = re.sub(r"(?i)<br\s*/?>|</(p|div|li|tr|h\d)>", "\n", s)
    return re.sub(r"\n\s*\n+", "\n\n", re.sub(r"<[^>]+>", " ", s))


def _load_structured(text: str, ext: str):
    if ext == ".json":
        return json.loads(text)
    import yaml  # noqa: PLC0415
    return yaml.safe_load(text)


def _largest_object_list(node):
    """The longest list of objects. That list is the table of example entries."""
    best_path, best = "", None

    def walk(cur, path):
        nonlocal best_path, best
        if isinstance(cur, list):
            if cur and isinstance(cur[0], dict) and (best is None or len(cur) > len(best)):
                best_path, best = path or "root", cur
            for item in cur:
                if isinstance(item, (dict, list)):
                    walk(item, path)
        elif isinstance(cur, dict):
            for key, value in cur.items():
                if isinstance(value, (dict, list)):
                    walk(value, f"{path}.{key}" if path else str(key))

    if isinstance(node, (dict, list)):
        walk(node, "")
    return (best_path, best) if best is not None else None


def _gather_services(node, found: set) -> None:
    if isinstance(node, dict):
        services = node.get("services")
        if isinstance(services, dict):
            found.update(str(k) for k in services)
        for value in node.values():
            if isinstance(value, (dict, list)):
                _gather_services(value, found)
    elif isinstance(node, list):
        for item in node:
            if isinstance(item, (dict, list)):
                _gather_services(item, found)


def _short(value) -> str:
    if isinstance(value, (str, int, float, bool)) or value is None:
        shown = json.dumps(value, ensure_ascii=False)
        return shown if len(shown) <= 80 else shown[:77] + "..."
    if isinstance(value, dict):
        names = [str(k) for k in list(value)[:12]]
        more = f", … +{len(value) - 12}" if len(value) > 12 else ""
        return f"[{', '.join(names)}{more}] ({len(value)} keys)"
    if isinstance(value, list):
        return f"list of {len(value)}"
    return type(value).__name__


def _example_line(item) -> str:
    if not isinstance(item, dict):
        return _short(item)
    bits = [f"{key}={_short(value)}" for key, value in item.items()]
    line = "; ".join(bits)
    return line if len(line) <= 500 else line[:497] + "..."


def summarize_structured(text: str, ext: str) -> str:
    """Replace a large JSON or YAML file with keys, counts, services, and a few examples."""
    note = (
        f"{SUMMARY_MARK} The full file is test data for the mock and is not included here."
    )
    lines = [note, f"original_chars: {len(text)}"]
    try:
        data = _load_structured(text, ext)
    except Exception as e:  # noqa: BLE001
        lines.append(f"parse: failed ({type(e).__name__})")
        return "\n".join(lines) + "\n"
    if isinstance(data, dict):
        keys = [str(k) for k in data]
        shown = keys if len(keys) <= 40 else keys[:40] + [f"… ({len(keys)} keys)"]
        lines.append("keys: " + ", ".join(shown))
        counts = []
        for key, value in data.items():
            path = str(key)
            if isinstance(value, list):
                counts.append(f"{path}: {len(value)}")
            elif isinstance(value, dict):
                counts.append(f"{path}: {len(value)} keys")
            elif isinstance(value, str) and len(value) <= 2000:
                counts.append(f"{path}: {value}")
            elif isinstance(value, str):
                counts.append(f"{path}: text, {len(value)} chars")
            else:
                counts.append(f"{path}: {_short(value)}")
        if len(counts) > 40:
            counts = counts[:40] + [f"… {len(counts) - 40} more keys"]
        lines.extend(counts)
    elif isinstance(data, list):
        lines.append("keys: (root list)")
        lines.append(f"count: {len(data)}")
    else:
        lines.append(f"value: {type(data).__name__}")
    services: set[str] = set()
    if isinstance(data, (dict, list)):
        _gather_services(data, services)
    lines.append("services: " + (", ".join(sorted(services)) if services else "(none)"))
    picked = _largest_object_list(data)
    if picked:
        path, rows = picked
        sample = rows[:3]
        lines.append(f"examples: {len(sample)} of {len(rows)} from {path}")
        for i, row in enumerate(sample, 1):
            lines.append(f"{i}. {_example_line(row)}")
    return "\n".join(lines) + "\n"


def to_text(p: Path) -> str | None:
    ext = p.suffix.lower()
    if ext in SKIP or ext in IMAGE or p.name.startswith("."):
        return None
    if ext in TEXT or ext in (".json", ".yaml", ".yml"):
        text = p.read_text(encoding="utf-8", errors="replace")
        if ext in (".json", ".yaml", ".yml") and len(text) > TEXT_CAP:
            return summarize_structured(text, ext)
        return text
    if ext == ".docx":
        return read_docx(p)
    if ext == ".pdf":
        return read_pdf(p)
    if ext in (".xlsx", ".xlsm"):
        return read_xlsx(p)
    if ext == ".csv":
        return read_csv(p)
    if ext in (".html", ".htm"):
        return read_html(p)
    return None


def _openapi_of(p: Path):
    try:
        if p.suffix.lower() == ".json":
            d = json.loads(p.read_text(encoding="utf-8"))
        elif p.suffix.lower() in (".yaml", ".yml"):
            import yaml  # noqa: PLC0415
            d = yaml.safe_load(p.read_text(encoding="utf-8"))
        else:
            return None
    except Exception:  # noqa: BLE001
        return None
    return d if isinstance(d, dict) and "openapi" in d else None


def fetch_openapi(mock_url: str):
    req = urllib.request.Request(mock_url.rstrip("/") + "/openapi.json", headers={"user-agent": "Mozilla/5.0 (builder)", "accept": "application/json"})
    with urllib.request.urlopen(req, timeout=30) as r:
        d = json.load(r)
    return d if isinstance(d, dict) and "openapi" in d and d.get("paths") else None   # a spec without paths is no tool API


def ingest(folder: Path, mock_url: str, out: Path) -> dict:
    folder, out = Path(folder), Path(out)
    out.mkdir(parents=True, exist_ok=True)
    skip_dirs = {".git", "node_modules", "vendor", "dist", "__pycache__"}
    files = sorted((p for p in folder.rglob("*") if p.is_file() and not any(
        part.startswith(".") or part in skip_dirs for part in p.relative_to(folder).parts)),
                   key=lambda p: (0 if p.name.lower() == FIRST[0] else 1 if p.name.lower() == FIRST[1] else 2, str(p.relative_to(folder)).lower()))
    from builder.ledger import describe_source  # noqa: PLC0415

    notes, parts, sources, api, api_src = [], [], [], None, None
    for p in files:
        rel = str(p.relative_to(folder))
        oa = _openapi_of(p)
        if oa is not None:
            api, api_src = api or oa, api_src or rel
            notes.append(f"{rel}: OpenAPI file (tool API fallback), not part of the brief text")
            continue
        if p.suffix.lower() in IMAGE:
            meta = describe_source(rel, "")
            meta["type"] = "drawing"
            meta["text"] = ""
            sources.append(meta)
            notes.append(f"{rel}: drawing, recorded, not read as text")
            continue
        try:
            t = to_text(p)
        except Exception as e:  # noqa: BLE001
            notes.append(f"{rel}: CONVERSION FAILED ({type(e).__name__}: {str(e)[:80]})")
            continue
        if t is None or not t.strip():
            notes.append(f"{rel}: skipped (binary or empty)")
            continue
        meta = describe_source(rel, t)
        meta["text"] = t
        sources.append(meta)
        date = meta.get("date") or "unknown"
        author = (meta.get("author") or "unknown").replace("\n", " ")
        parts.append(f"=== FILE: {rel} ===\n=== META: date={date}; type={meta['type']}; author={author} ===\n\n{t.strip()}\n")
        kind = "summarized" if t.startswith(SUMMARY_MARK) else f"{len(t)} chars"
        notes.append(f"{rel}: {kind}; date={date}; type={meta['type']}; author={author}")
    try:
        live = fetch_openapi(mock_url) if mock_url else None
    except Exception as e:  # noqa: BLE001
        live = None
        notes.append(f"GET {mock_url}/openapi.json failed ({type(e).__name__}); looking in the folder")
    if live is not None:
        api, api_src = live, f"{mock_url.rstrip('/')}/openapi.json"
    if api is None:
        raise RuntimeError("no tool API: {mock_url}/openapi.json is unavailable and the folder has no OpenAPI file")
    servers = [s for s in api.get("servers") or [] if isinstance(s, dict) and str(s.get("url", "")).startswith("http")]
    if not servers and mock_url:   # no usable server URL in the spec: tools go to the mock URL
        api["servers"] = [{"url": mock_url.rstrip("/")}]
        notes.append(f"openapi has no absolute servers[].url: set to {mock_url.rstrip('/')}")
    brief = out / "case_brief.md"
    brief.write_text("\n".join(parts), encoding="utf-8")
    api_path = out / "tools.openapi.json"
    api_path.write_text(json.dumps(api, ensure_ascii=False, indent=1), encoding="utf-8")
    (out / "ingest_notes.txt").write_text("\n".join(notes) + f"\ntool API from: {api_src}\n", encoding="utf-8")
    return {"brief": str(brief), "api": str(api_path), "notes": notes, "api_source": api_src,
            "files": len(parts), "chars": brief.stat().st_size, "sources": sources}


if __name__ == "__main__":
    r = ingest(Path(sys.argv[1]), sys.argv[2], Path(sys.argv[3]))
    print(json.dumps({k: r[k] for k in ("brief", "api", "files", "chars", "api_source")}))
    print("\n".join(r["notes"]))
