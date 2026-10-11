import json
import sys
from pathlib import Path

REPO = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(REPO))


def test_ingest_order_headers_and_api_fallback(tmp_path, monkeypatch):
    import docx, openpyxl
    from builder import ingest as I
    c = tmp_path / "case"; c.mkdir()
    (c / "BRIEF.md").write_text("ana kural şğü"); (c / "README.md").write_text("readme")
    (c / "aaa_sss.md").write_text("soru cevap")
    d = docx.Document(); d.add_paragraph("duyuru metni"); t = d.add_table(rows=1, cols=2); t.rows[0].cells[0].text = "a"; t.rows[0].cells[1].text = "b"; d.save(str(c / "duyuru.docx"))
    wb = openpyxl.Workbook(); wb.active.append(["x", "y"]); wb.active.append([1, 2]); wb.save(str(c / "t.xlsx"))
    (c / "spec.json").write_text(json.dumps({"openapi": "3.0.0", "paths": {"/a": {}}, "servers": [{"url": "https://real.example/v1"}]}))
    monkeypatch.setattr(I, "fetch_openapi", lambda u: (_ for _ in ()).throw(OSError("down")))
    r = I.ingest(c, "http://mock.example", tmp_path / "out")
    text = Path(r["brief"]).read_text()
    heads = [l for l in text.splitlines() if l.startswith("=== FILE:")]
    assert heads[:2] == ["=== FILE: BRIEF.md ===", "=== FILE: README.md ==="] and len(heads) == 5
    assert "duyuru metni" in text and "| x | y |" in text and "ana kural şğü" in text and "spec.json" not in text
    assert json.loads(Path(r["api"]).read_text())["servers"][0]["url"] == "https://real.example/v1"


def _big_customers(n=40, blob=800):
    customers = []
    for i in range(n):
        customers.append({
            "customerNo": f"{i:08d}",
            "firstName": "Ada" if i == 0 else f"Name{i}",
            "services": {
                "cards.list": {"http": 200, "response": {"blob": "X" * blob, "mark": f"only-in-body-{i}"}},
                "cards.balance": {"http": 200, "response": {"blob": "Y" * blob}},
            },
        })
    return {"description": "synthetic ledger", "customers": customers}


def test_large_json_is_summarized_and_small_json_is_kept(tmp_path, monkeypatch):
    from builder import ingest as I
    case = tmp_path / "case"
    case.mkdir()
    (case / "BRIEF.md").write_text("The assistant must greet the caller.")
    (case / "rows.json").write_text('{"rows": [{"id": "keep-me"}]}')
    big = json.dumps(_big_customers(), ensure_ascii=False)
    assert len(big) > I.TEXT_CAP
    (case / "customers.json").write_text(big)
    (case / "api.json").write_text(json.dumps({"openapi": "3.0.0", "paths": {"/a": {}}, "servers": [{"url": "https://tools.example/v1"}]}))
    monkeypatch.setattr(I, "fetch_openapi", lambda u: (_ for _ in ()).throw(OSError("down")))
    result = I.ingest(case, "http://mock.example", tmp_path / "out")
    text = Path(result["brief"]).read_text()
    assert '"keep-me"' in text
    assert "only-in-body-0" not in text and "only-in-body-9" not in text
    assert "X" * 800 not in text
    assert "The full file is test data for the mock" in text
    assert "keys: description, customers" in text
    assert "customers: 40" in text
    assert "services: cards.balance, cards.list" in text
    assert "examples: 3 of 40 from customers" in text
    assert 'customerNo="00000000"' in text and 'customerNo="00000002"' in text
    assert 'customerNo="00000003"' not in text
    assert len(text) < I.TEXT_CAP
    notes = (tmp_path / "out" / "ingest_notes.txt").read_text()
    assert "customers.json: summarized" in notes
    row_line = next(line for line in notes.splitlines() if line.startswith("rows.json:"))
    assert "summarized" not in row_line


def test_real_case_folder_mails_stay_and_big_json_is_summarized(tmp_path, monkeypatch):
    import re
    from builder import ingest as I
    case = Path("/Users/toprakgundogdu/Downloads/test-environment")
    ledger = case / "test-customers.json"
    if not ledger.is_file():
        import pytest
        pytest.skip("test-environment is not beside the repo")
    doc = json.loads(ledger.read_text(encoding="utf-8"))
    early = json.dumps(doc["customers"][:3], ensure_ascii=False)
    mark = None
    for customer in doc["customers"][3:]:
        found = re.findall(r'"guid": "([0-9a-f]{16,})"', json.dumps(customer, ensure_ascii=False))
        if found and found[0] not in early:
            mark = found[0]
            break
    assert mark
    monkeypatch.setattr(I, "fetch_openapi", lambda u: {
        "openapi": "3.0.0", "paths": {"/a": {}}, "servers": [{"url": "http://127.0.0.1:8788"}],
    })
    result = I.ingest(case, "http://127.0.0.1:8788", tmp_path / "out")
    text = Path(result["brief"]).read_text(encoding="utf-8")
    heads = [line for line in text.splitlines() if line.startswith("=== FILE:")]
    assert heads[0] == "=== FILE: BRIEF.md ==="
    assert "=== FILE: sources/emails/2026-09-24-kickoff-questions.txt ===" in text
    assert "=== FILE: sources/emails/2026-10-09-decisions-recap.txt ===" in text
    assert "=== FILE: sources/emails/attachments/2026-09-02-statement-debt-analysis.md ===" in text
    assert "=== FILE: sources/documents/2026-09-10-ivr-prompt-sheet.md ===" in text
    assert "=== FILE: sources/documents/2025-08-14-cards-faq.md ===" in text
    assert "=== FILE: sources/meetings/2026-10-07-analysis-session-transcript.txt ===" in text
    assert "Please don't ask for one." in text
    assert "Three calming replies it is" in text
    assert "first pass after the two walkthroughs" in text
    assert "my camera is off" in text
    assert "DBT-010" in text
    assert "statement-mock" in text
    start = text.index("=== FILE: test-customers.json ===")
    nxt = text.find("\n=== FILE:", start + 1)
    section = text[start:] if nxt < 0 else text[start:nxt]
    assert "The full file is test data for the mock" in section
    assert "customers: 34" in section
    assert "examples: 3 of 34 from customers" in section
    assert 'customerNo="90000001"' in section
    for name in (
        "cards.balance", "cards.list", "contact.email.get", "statement.pdf.send",
        "statement.week.change", "statement.week.decide", "statement.week.get",
    ):
        assert name in section.split("services: ", 1)[1].splitlines()[0]
    assert mark not in section
    assert len(section) < I.TEXT_CAP
    assert ledger.stat().st_size > 700_000
    assert len(text) < ledger.stat().st_size
    assert "test-customers.json: summarized" in (tmp_path / "out" / "ingest_notes.txt").read_text()
